from pathlib import Path

from openpyxl import Workbook, load_workbook
import pytest

from app.report_writer import (
    write_area_report,
    write_overall_report,
    write_store_report,
)


@pytest.fixture
def aggregation_result():
    return {
        "sales_amount_tax_ex": 1_000_000,
        "tax_amount": 100_000,
        "sales_amount_tax_in": 1_100_000,
        "customer_count": 500,
        "average_customer_spend": 2_200,
        "tax_rate": 0.1,
        "cost_amount": 300_000,
        "labor_cost_amount": 250_000,
        "rent_cost": 100_000,
        "utility_cost": 50_000,
        "other_expense_cost": 30_000,
        "operating_cost": 430_000,
        "operating_profit": 270_000,
        "cost_rate": 0.3,
        "labor_cost_rate": 0.25,
        "operating_profit_rate": 0.27,
    }


@pytest.fixture
def template_path(tmp_path):
    path = tmp_path / "template.xlsx"
    Workbook().save(path)
    return path


@pytest.mark.parametrize("report_type", ["store", "area", "overall"])
def test_report_writer_saves_decimal_rates_with_percentage_format(
    report_type,
    aggregation_result,
    template_path,
    tmp_path,
):
    output_path = tmp_path / f"{report_type}.xlsx"
    write_report(
        report_type,
        aggregation_result,
        template_path,
        output_path,
    )

    worksheet = load_workbook(output_path).active

    assert worksheet["D18"].value == pytest.approx(0.3)
    assert worksheet["F18"].value == pytest.approx(0.25)
    assert worksheet["B21"].value == pytest.approx(0.27)
    assert worksheet["D18"].number_format == "0.0%"
    assert worksheet["F18"].number_format == "0.0%"
    assert worksheet["B21"].number_format == "0.0%"


def write_report(
    report_type: str,
    aggregation_result: dict,
    template_path: Path,
    output_path: Path,
) -> None:
    if report_type == "store":
        write_store_report(
            aggregation_result,
            "202606",
            "001",
            "店舗A",
            template_path,
            output_path,
        )
    elif report_type == "area":
        write_area_report(
            aggregation_result,
            "202606",
            "01",
            "エリアA",
            template_path,
            output_path,
        )
    else:
        write_overall_report(
            aggregation_result,
            "202606",
            template_path,
            output_path,
        )
